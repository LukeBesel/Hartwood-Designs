"""Picture-based instruction booklet for a kit, as PDF (reportlab) and Word (python-docx).

    python engine/booklet.py blaster [--no-render]

1. Renders one image per step (built parts gray, new parts in color with an orange outline, arrows).
2. Adds code callouts to each image (code and quantity, with leader lines).
3. Lays out: cover, before you start, what is in the box, parts maps, steps (parts tray, image, icons, tip,
   test, QR code to the same step in the 3D app), finish page.
Output: kits/<kit>/booklet/<kit>_instructions.pdf and .docx (+ images in booklet/img/).
"""
from __future__ import annotations

import io
import json
import math
import os
import re
import subprocess
import sys
import tempfile

import qrcode
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from hartwood import config as C  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
BRAND = os.path.join(ROOT, 'brand')
FONTS = os.path.join(BRAND, 'fonts')
ICONS = os.path.join(BRAND, 'icons')
PAL = {k: v['hex'] for k, v in json.load(open(os.path.join(BRAND, 'brand.json')))['palette'].items()}
for name, f in (('Display', 'Fraunces-700'), ('DisplayX', 'Fraunces-800'), ('Body', 'Inter-400'),
                ('BodyM', 'Inter-500'), ('BodyB', 'Inter-700')):
    pdfmetrics.registerFont(TTFont(name, os.path.join(FONTS, f + '.ttf')))
PW, PH = landscape(letter)          # 792 x 612 pt
M = 28

AGE = '14+'
RULES = [
    ('push-out', 'Push parts out from the back of the sheet. Do not twist them out.'),
    ('sort', 'Every part has a code. B12 is part 12 on sheet B. Codes for show faces are printed next to the part.'),
    ('press', 'Press tabs straight in, firm but gentle. No glue needed.'),
    ('sand', 'Too tight? Rub the tab on the sandpaper two or three times, then try again.'),
    ('wax', 'Rub wax into every hole that turns. It makes moving parts smooth.'),
    ('test', 'A check mark means: test it now, before you close it up.'),
    ('careful', f'Small parts. Not for children under {AGE[:-1]}. Keep spare parts until you finish.'),
]
BLASTER_RULES = [('glasses', 'Wear glasses when you fire. Only aim at the target. Never at faces, people or pets.')]


def font(name, size):
    f = {'Display': 'Fraunces-700', 'Body': 'Inter-500', 'BodyB': 'Inter-700'}[name]
    return ImageFont.truetype(os.path.join(FONTS, f + '.ttf'), size)


# ------------------------------------------------------------------ 1 render step images
def render_steps(kit, steps, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    jobs = []
    for s in steps['steps']:
        jobs.append({'glb': f'/kits/{kit}/{kit}.glb', 'steps': f'/kits/{kit}/{kit}_steps.json',
                     'out': os.path.relpath(os.path.join(out_dir, f"step_{s['n']:02d}.png"), ROOT), 'mode': 'step',
                     'step': s['n'], 'view': s['view'] or 'auto', 'explode_new': 0.7, 'width': 1600, 'height': 1150,
                     'margin': 1.08})
    with tempfile.TemporaryDirectory() as td:
        jp, mp = os.path.join(td, 'jobs.json'), os.path.join(out_dir, 'steps_meta.json')
        json.dump(jobs, open(jp, 'w'))
        subprocess.run(['node', os.path.join(ROOT, 'engine/render/render.mjs'), jp, mp], check=True, cwd=ROOT,
                       stdout=subprocess.DEVNULL)


# ------------------------------------------------------------------ 2 callouts
def label_text(lab):
    if lab['kind'] != 'part':
        return 'Dowel'
    return lab['code'] + (f"  x{lab['qty']}" if lab['qty'] > 1 else '')


def annotate(png, labels, out):
    im = Image.open(png).convert('RGBA')
    W, H = im.size
    bg = Image.new('RGBA', im.size, (255, 255, 255, 255))
    bg.alpha_composite(im)
    d = ImageDraw.Draw(bg)
    f = font('BodyB', 34)
    cx, cy = W / 2, H / 2
    placed = []
    accent = tuple(int(PAL['blaze'][i:i + 2], 16) for i in (1, 3, 5))
    ink = tuple(int(PAL['walnut'][i:i + 2], 16) for i in (1, 3, 5))
    for lab in sorted(labels, key=lambda l: (l['pts'][0][1], l['pts'][0][0])):
        pts = lab['pts']
        ax = sum(p[0] for p in pts) / len(pts); ay = sum(p[1] for p in pts) / len(pts)
        txt = label_text(lab)
        tw = d.textlength(txt, font=f); bw, bh = tw + 32, 54
        base = math.atan2(ay - cy, ax - cx) if (ax, ay) != (cx, cy) else -math.pi / 2
        best = None
        for dist in (150, 210, 270, 340, 420):
            for da in (0, 0.35, -0.35, 0.7, -0.7, 1.05, -1.05, 1.4, -1.4, 1.75, -1.75, math.pi):
                a = base + da
                bx = ax + dist * math.cos(a) - bw / 2; by = ay + dist * math.sin(a) - bh / 2
                bx = min(max(bx, 10), W - bw - 10); by = min(max(by, 10), H - bh - 10)
                rect = (bx - 8, by - 8, bx + bw + 8, by + bh + 8)
                if any(not (rect[2] < r[0] or rect[0] > r[2] or rect[3] < r[1] or rect[1] > r[3]) for r in placed):
                    continue
                best = (bx, by); break
            if best:
                break
        if best is None:
            best = (min(max(ax + 60, 10), W - bw - 10), min(max(ay - 80, 10), H - bh - 10))
        bx, by = best
        placed.append((bx, by, bx + bw, by + bh))
        bxc, byc = bx + bw / 2, by + bh / 2
        near = min(pts, key=lambda p: (p[0] - bxc) ** 2 + (p[1] - byc) ** 2)
        d.line([(bxc, byc), tuple(near)], fill=ink, width=3)
        for (px, py) in pts[:12]:          # one leader line; the other copies get a dot
            d.ellipse([px - 7, py - 7, px + 7, py + 7], fill=accent, outline=ink, width=2)
    for (bx, by, bx2, by2), lab in zip(placed, sorted(labels, key=lambda l: (l['pts'][0][1], l['pts'][0][0]))):
        d.rounded_rectangle([bx, by, bx2, by2], 14, fill=(255, 255, 255), outline=accent, width=4)
        d.text((bx + 16, by + 8), label_text(lab), font=f, fill=ink)
    bg.convert('RGB').save(out, quality=88)


# ------------------------------------------------------------------ helpers
def qr_png(url, path):
    q = qrcode.QRCode(border=1, box_size=10, error_correction=qrcode.constants.ERROR_CORRECT_M)
    q.add_data(url); q.make(fit=True)
    q.make_image(fill_color=PAL['walnut'], back_color='white').save(path)
    return path


def step_url(kit, n):
    return f'{C.SITE_URL}/build/{kit}/step/{n}'


def wrap(c, text, fontname, size, width):
    words, lines, cur = text.split(), [], ''
    for w in words:
        t = (cur + ' ' + w).strip()
        if pdfmetrics.stringWidth(t, fontname, size) > width and cur:
            lines.append(cur); cur = w
        else:
            cur = t
    if cur:
        lines.append(cur)
    return lines


def draw_text(c, x, y, text, fontname='Body', size=12, width=400, leading=None, color=None):
    c.setFont(fontname, size); c.setFillColor(HexColor(color or PAL['walnut']))
    leading = leading or size * 1.3
    for line in wrap(c, text, fontname, size, width):
        c.drawString(x, y, line); y -= leading
    return y


def fit_image(c, path, x, y, w, h, align='center'):
    ir = ImageReader(path)
    iw, ih = ir.getSize()
    s = min(w / iw, h / ih)
    dw, dh = iw * s, ih * s
    ox = x + (w - dw) / 2 if align == 'center' else x
    c.drawImage(ir, ox, y + (h - dh) / 2, dw, dh, mask='auto')
    return dw, dh


def page_bg(c):
    c.setFillColor(HexColor(PAL['paper'])); c.rect(0, 0, PW, PH, stroke=0, fill=1)


def footer(c, kit_name, page):
    c.setFont('Body', 9); c.setFillColor(HexColor(PAL['oak']))
    c.drawString(M, 14, f'Hartwood Designs  |  {kit_name}')
    c.drawRightString(PW - M, 14, str(page))


def part_thumb(kit, code):
    return os.path.join(ROOT, 'kits', kit, 'parts', f'{code}.png')


# ------------------------------------------------------------------ 3 PDF
def build_pdf(kit, steps, report, img_dir, out_pdf):
    name = steps['name']
    c = canvas.Canvas(out_pdf, pagesize=(PW, PH))
    c.setTitle(f'{name} instructions'); c.setAuthor('Hartwood Designs')
    page = [1]

    def new_page():
        footer(c, name, page[0]); c.showPage(); page[0] += 1; page_bg(c)

    # cover
    page_bg(c)
    c.setFillColor(HexColor(PAL['birch'])); c.rect(0, 0, PW, 150, stroke=0, fill=1)
    fit_image(c, os.path.join(BRAND, 'logo.png'), M, PH - 70, 210, 46, align='left')
    fit_image(c, os.path.join(ROOT, 'kits', kit, 'renders', 'three_quarter.jpg'), 300, 120, PW - 300 - M, PH - 150)
    c.setFont('DisplayX', 50); c.setFillColor(HexColor(PAL['walnut'])); c.drawString(M, PH - 160, name)
    y = draw_text(c, M, PH - 196, steps['tagline'], 'BodyM', 15, 260, color=PAL['lagoon_dark'])
    badges = [f"{steps['piece_count']} pieces", f"About {steps['build_hours']}", f'Ages {AGE}', 'No glue needed',
              f"{len(steps['steps'])} steps"]
    by = 100
    for i, b in enumerate(badges):
        bx = M + (i % 3) * 140; yy = by - (i // 3) * 36
        c.setFillColor(HexColor(PAL['lagoon'])); c.roundRect(bx, yy, 130, 26, 8, stroke=0, fill=1)
        c.setFillColor(white); c.setFont('BodyB', 11); c.drawCentredString(bx + 65, yy + 8.5, b)
    new_page()

    # before you start
    c.setFont('Display', 30); c.setFillColor(HexColor(PAL['walnut'])); c.drawString(M, PH - 60, 'Before you start')
    rules = RULES + (BLASTER_RULES if kit == 'blaster' else [])
    y = PH - 110
    for icon, text in rules:
        fit_image(c, os.path.join(ICONS, icon + '.png'), M, y - 14, 40, 40)
        draw_text(c, M + 52, y + 8, text, 'BodyM', 13, 330)
        y -= 56
    # what's in the box
    x2 = 440
    c.setFont('Display', 20); c.setFillColor(HexColor(PAL['walnut'])); c.drawString(x2, PH - 104, "What's in the box")
    yy = PH - 130
    for sh in steps['sheets']:
        draw_text(c, x2, yy, f"Sheet {sh['letter']}: {sh['label']}, {sh['parts']} parts"
                  + (f" + {sh['spares']} spares" if sh['spares'] else ''), 'Body', 11.5, 320); yy -= 18
    yy -= 8
    for k, v in steps['extras'].items():
        yy = draw_text(c, x2, yy, f'{v} x {k}', 'Body', 11.5, 320) - 2
    q = qr_png(f'{C.SITE_URL}/build/{kit}', os.path.join(img_dir, 'qr_build.png'))
    fit_image(c, q, x2, 60, 80, 80, align='left')
    draw_text(c, x2 + 92, 118, 'Build it in 3D on your phone: scan this code. Turn the model, tap any part, '
              'and watch it slide into place.', 'BodyM', 11.5, 220, color=PAL['lagoon_dark'])
    new_page()

    # parts maps (two per page)
    maps = sorted(os.listdir(os.path.join(ROOT, 'kits', kit, 'parts_map')))
    for i in range(0, len(maps), 2):
        c.setFont('Display', 22); c.setFillColor(HexColor(PAL['walnut'])); c.drawString(M, PH - 52, 'Parts maps')
        draw_text(c, M + 160, PH - 50, 'Find each part by its sheet letter and number.', 'Body', 11, 400)
        for j, mfile in enumerate(maps[i:i + 2]):
            fit_image(c, os.path.join(ROOT, 'kits', kit, 'parts_map', mfile), M + j * ((PW - 2 * M) / 2 + 6), 40,
                      (PW - 2 * M) / 2 - 6, PH - 110)
        new_page()

    # steps
    for s in steps['steps']:
        n = s['n']
        c.setFillColor(HexColor(PAL['lagoon'])); c.circle(M + 26, PH - 50, 26, stroke=0, fill=1)
        c.setFillColor(white); c.setFont('DisplayX', 26); c.drawCentredString(M + 26, PH - 59, str(n))
        c.setFillColor(HexColor(PAL['walnut'])); c.setFont('Display', 21)
        tl = wrap(c, s['title'], 'Display', 21, PW - 2 * M - 160)
        for k, line in enumerate(tl[:2]):
            c.drawString(M + 66, PH - 46 - k * 24 + (10 if len(tl) > 1 else 0) - 8, line)
        q = qr_png(step_url(kit, n), os.path.join(img_dir, f'qr_{n:02d}.png'))
        fit_image(c, q, PW - M - 58, PH - 84, 58, 58)
        c.setFont('Body', 8); c.setFillColor(HexColor(PAL['lagoon_dark']))
        c.drawRightString(PW - M - 64, PH - 76, 'See it in 3D')
        # parts tray
        tray_x, tray_w, top = M, 176, PH - 96
        c.setFillColor(HexColor(PAL['birch'])); c.roundRect(tray_x, 90, tray_w, top - 90, 10, stroke=0, fill=1)
        tiles = [(p['code'], p['qty'], part_thumb(kit, p['code'])) for p in s['parts']]
        for h in s['hardware']:
            mm = re.search(r'([0-9.]+) mm', h['note'] or '')
            tiles.append(('Dowel' + (f" {mm.group(1)} mm" if mm else ''), 1, None))
        cols = 2 if len(tiles) <= 10 else 3
        tw = (tray_w - 12) / cols; th = tw * 0.95 if cols == 2 else tw * 1.05
        rows = math.ceil(len(tiles) / cols) if tiles else 0
        if rows and rows * th > top - 100:
            th = (top - 100) / rows
        for k, (code, qty, img) in enumerate(tiles):
            tx = tray_x + 6 + (k % cols) * tw; ty = top - 8 - (k // cols + 1) * th
            if img and os.path.exists(img):
                fit_image(c, img, tx + 2, ty + 14, tw - 4, th - 18)
            elif code.startswith('Dowel'):
                c.setStrokeColor(HexColor(PAL['oak'])); c.setLineWidth(5); c.setLineCap(1)
                c.line(tx + 10, ty + th / 2 + 4, tx + tw - 10, ty + th / 2 + 4)
            c.setFillColor(HexColor(PAL['walnut'])); c.setFont('BodyB', 10 if cols == 2 else 8.5)
            c.drawCentredString(tx + tw / 2, ty + 3, code + (f'  x{qty}' if qty > 1 else ''))
        if not tiles:
            draw_text(c, tray_x + 14, top - 30, 'No new parts in this step.', 'BodyM', 11, tray_w - 28)
        # main image
        ann = os.path.join(img_dir, f'step_{n:02d}.jpg')
        fit_image(c, ann, tray_x + tray_w + 10, 90, PW - M - (tray_x + tray_w + 10), top - 90)
        # bottom band: icons, tip, test
        x = M
        for icon in s['icons'] + (['test'] if s['test'] else []):
            p = os.path.join(ICONS, icon + '.png')
            if os.path.exists(p):
                fit_image(c, p, x, 34, 44, 44, align='left'); x += 50
        tip_w = (PW - M - x - 12) * (0.5 if s['test'] else 1.0)
        yb = draw_text(c, x + 6, 66, s['tip'], 'BodyM', 12, tip_w)
        if s['test']:
            tx0 = x + 18 + tip_w
            c.setFillColor(HexColor(PAL['lagoon_light'])); c.roundRect(tx0, 26, PW - M - tx0, 58, 8, stroke=0, fill=1)
            c.setFillColor(HexColor(PAL['lagoon_dark'])); c.setFont('BodyB', 11); c.drawString(tx0 + 10, 70, 'TEST IT NOW')
            draw_text(c, tx0 + 10, 55, s['test'], 'Body', 10.5, PW - M - tx0 - 20, color=PAL['lagoon_dark'])
        new_page()

    # finish
    fit_image(c, os.path.join(ROOT, 'kits', kit, 'renders', 'front.jpg'), 330, 60, PW - 330 - M, PH - 100)
    c.setFont('DisplayX', 40); c.setFillColor(HexColor(PAL['walnut'])); c.drawString(M, PH - 90, 'You built it!')
    y = PH - 130
    finish = ['Dust it with a soft dry brush. Keep it out of hot sun and damp rooms.',
              'If a part works loose, press it back in. A tiny drop of wood glue is fine for repairs.',
              'Missing or broken part? Send us the code from the parts map and we will mail a new one.']
    if kit == 'blaster':
        finish = ['Only fire #64 rubber bands. Load at most 7.', 'Wear glasses. Aim only at the target. '
                  'Never at faces, people or pets.', 'Keep the orange nose, plug and cap on the ship. They show it is a toy.'] + finish
    for t in finish:
        y = draw_text(c, M, y, t, 'BodyM', 13, 280) - 10
    draw_text(c, M, 120, 'Help and spare parts: TODO_LUCAS support email', 'Body', 11, 280, color=PAL['lagoon_dark'])
    q = qr_png(f'{C.SITE_URL}/kits/{kit}', os.path.join(img_dir, 'qr_kit.png'))
    fit_image(c, q, M, 36, 70, 70, align='left')
    draw_text(c, M + 80, 80, 'Share a photo of your build and find our other kits.', 'Body', 11, 200)
    footer(c, name, page[0]); c.showPage()
    c.save()
    return page[0]


# ------------------------------------------------------------------ 4 Word
def build_docx(kit, steps, img_dir, out_docx):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Inches(0.5))
    st = doc.styles['Normal']; st.font.name = 'Inter'; st.font.size = Pt(11)
    walnut = RGBColor.from_string(PAL['walnut'][1:]); lag = RGBColor.from_string(PAL['lagoon_dark'][1:])

    def heading(text, size=26):
        p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(size)
        r.font.name = 'Fraunces'; r.font.color.rgb = walnut
        return p

    def jpg(path, maxw=None):
        im = Image.open(path)
        if im.mode in ('RGBA', 'LA', 'P'):
            bg = Image.new('RGB', im.size, (255, 255, 255)); im = im.convert('RGBA'); bg.paste(im, mask=im.split()[3]); im = bg
        if maxw and im.width > maxw:
            im = im.resize((maxw, int(im.height * maxw / im.width)))
        b = io.BytesIO(); im.save(b, 'JPEG', quality=85); b.seek(0)
        return b
    doc.add_picture(os.path.join(BRAND, 'logo.png'), width=Inches(2.6))
    heading(steps['name'], 40)
    p = doc.add_paragraph(steps['tagline']); p.runs[0].font.color.rgb = lag; p.runs[0].font.size = Pt(14)
    doc.add_paragraph(f"{steps['piece_count']} pieces  |  About {steps['build_hours']}  |  Ages {AGE}  |  No glue needed")
    doc.add_picture(jpg(os.path.join(ROOT, 'kits', kit, 'renders', 'three_quarter.jpg'), 1600), width=Inches(7.5))
    doc.add_page_break()
    heading('Before you start')
    t = doc.add_table(rows=0, cols=2)
    for icon, text in RULES + (BLASTER_RULES if kit == 'blaster' else []):
        row = t.add_row().cells
        row[0].paragraphs[0].add_run().add_picture(os.path.join(ICONS, icon + '.png'), width=Inches(0.5))
        row[1].paragraphs[0].add_run(text).font.size = Pt(13)
        row[0].width = Inches(0.7); row[1].width = Inches(8.5)
    heading("What's in the box", 18)
    for sh in steps['sheets']:
        doc.add_paragraph(f"Sheet {sh['letter']}: {sh['label']}, {sh['parts']} parts" +
                          (f" + {sh['spares']} spares" if sh['spares'] else ''), style='List Bullet')
    for k, v in steps['extras'].items():
        doc.add_paragraph(f'{v} x {k}', style='List Bullet')
    doc.add_page_break()
    heading('Parts maps', 22)
    for mfile in sorted(os.listdir(os.path.join(ROOT, 'kits', kit, 'parts_map'))):
        doc.add_picture(jpg(os.path.join(ROOT, 'kits', kit, 'parts_map', mfile), 1400), width=Inches(4.6))
    for s in steps['steps']:
        n = s['n']
        doc.add_page_break()
        ht = doc.add_table(rows=1, cols=2)
        r = ht.rows[0].cells
        r[0].width = Inches(8.8); r[1].width = Inches(1.2)
        pr = r[0].paragraphs[0]
        rn = pr.add_run(f'Step {n}   '); rn.bold = True; rn.font.size = Pt(24); rn.font.color.rgb = lag
        rt = pr.add_run(s['title']); rt.bold = True; rt.font.size = Pt(20); rt.font.color.rgb = walnut
        qp = r[1].paragraphs[0]; qp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        qp.add_run().add_picture(os.path.join(img_dir, f'qr_{n:02d}.png'), width=Inches(0.9))
        bt = doc.add_table(rows=1, cols=2)
        left, right = bt.rows[0].cells
        left.width = Inches(2.2); right.width = Inches(7.8)
        lp = left.paragraphs[0]
        for p in s['parts']:
            img = part_thumb(kit, p['code'])
            if os.path.exists(img):
                lp.add_run().add_picture(jpg(img), width=Inches(0.95))
            lr = lp.add_run(f"  {p['code']}" + (f" x{p['qty']}" if p['qty'] > 1 else '') + '\n'); lr.bold = True
        for h in s['hardware']:
            lp.add_run('Dowel: ' + h['note'] + '\n')
        right.paragraphs[0].add_run().add_picture(jpg(os.path.join(img_dir, f'step_{n:02d}.jpg'), 1400), width=Inches(7.6))
        tp = doc.add_paragraph()
        for icon in s['icons'] + (['test'] if s['test'] else []):
            ip = os.path.join(ICONS, icon + '.png')
            if os.path.exists(ip):
                tp.add_run().add_picture(ip, width=Inches(0.42)); tp.add_run(' ')
        tr = tp.add_run('  ' + s['tip']); tr.font.size = Pt(13)
        if s['test']:
            q = doc.add_paragraph(); a = q.add_run('TEST IT NOW: '); a.bold = True; a.font.color.rgb = lag
            q.add_run(s['test']).font.color.rgb = lag
    doc.add_page_break()
    heading('You built it!', 34)
    doc.add_picture(jpg(os.path.join(ROOT, 'kits', kit, 'renders', 'front.jpg'), 1600), width=Inches(6.5))
    doc.add_paragraph('Help and spare parts: TODO_LUCAS support email')
    doc.core_properties.title = f"{steps['name']} instructions"; doc.core_properties.author = 'Hartwood Designs'
    doc.save(out_docx)


def main(kit, render=True):
    steps = json.load(open(os.path.join(ROOT, 'kits', kit, f'{kit}_steps.json')))
    report = json.load(open(os.path.join(ROOT, 'kits', kit, f'{kit}_report.json')))
    bdir = os.path.join(ROOT, 'kits', kit, 'booklet'); img_dir = os.path.join(bdir, 'img')
    raw = os.path.join(bdir, 'raw')
    os.makedirs(img_dir, exist_ok=True)
    if render or not os.path.exists(os.path.join(raw, 'steps_meta.json')):
        render_steps(kit, steps, raw)
    meta = {os.path.basename(m['out']): m for m in json.load(open(os.path.join(raw, 'steps_meta.json')))}
    for s in steps['steps']:
        f = f"step_{s['n']:02d}.png"
        annotate(os.path.join(raw, f), meta[f]['labels'], os.path.join(img_dir, f"step_{s['n']:02d}.jpg"))
    pages = build_pdf(kit, steps, report, img_dir, os.path.join(bdir, f'{kit}_instructions.pdf'))
    build_docx(kit, steps, img_dir, os.path.join(bdir, f'{kit}_instructions.docx'))
    # the 3D app reads the step views chosen here so phone and paper match
    views = {int(m['out'].split('_')[-1][:2]): m['view'] for m in meta.values()}
    json.dump(views, open(os.path.join(bdir, 'step_views.json'), 'w'))
    print(f'{kit}: booklet {pages} pages')


if __name__ == '__main__':
    main(sys.argv[1], render='--no-render' not in sys.argv)
